from __future__ import annotations

import hashlib
import io
from dataclasses import replace
from types import SimpleNamespace
from unittest.mock import Mock

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
from sqlalchemy import func, select

from app.api import documents as document_api
from app.core.config import get_settings
from app.core.database import get_db
from app.models.document import Document, DocumentChunk, DocumentVersion, DocumentVersionChunk
from app.services.indexing.evidence_resolver import resolve_version_chunk
from app.services.indexing.ingestion import DocumentIngestion, RawSourceUnavailable
from app.services.indexing.incremental_indexer import IncrementalIndexer, IncrementalIndexError
from app.services.indexing.reconciliation import VectorReconciliationService
from app.services.indexing.source_gc import SourceGarbageCollector
from app.services.indexing.version_chunks import version_chunks
from app.services.parsing.base import ParsedDocument, ParsedSection
from app.services.parsing.url_parser import UrlParser
from app.services.retrieval.retrieval_service import ResearchRetrievalService
from app.storage.local import LocalSourceStorage

pytestmark = pytest.mark.integration


@pytest.fixture
def evidence_context(integration_db, vector_store, rule_extractor, tmp_path):
    get_settings().source_storage_root = str(tmp_path / "sources")
    storage = LocalSourceStorage()
    indexer = IncrementalIndexer(vector_store, rule_extractor)
    return SimpleNamespace(db=integration_db, store=vector_store, storage=storage, indexer=indexer,
                           ingestion=DocumentIngestion(indexer, storage))


def pdf_bytes(*pages):
    writer = PdfWriter()
    for value in pages:
        page = writer.add_blank_page(width=612, height=792)
        font = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"),
                                 NameObject("/BaseFont"): NameObject("/Helvetica")})
        page[NameObject("/Resources")] = DictionaryObject({
            NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})})
        stream = DecodedStreamObject()
        stream.set_data(f"BT /F1 12 Tf 50 700 Td ({value}) Tj ET".encode("ascii"))
        page[NameObject("/Contents")] = writer._add_object(stream)
    output = io.BytesIO()
    writer.write(output)
    return output.getvalue()


def versions(ctx, document_id):
    return list(ctx.db.scalars(select(DocumentVersion).where(DocumentVersion.document_id == document_id)
                              .order_by(DocumentVersion.version)).all())


def text_import(ctx, text="Alpha evidence.\n\nBeta evidence.\n\nOld evidence.", uri="phase2:study"):
    return ctx.ingestion.import_text(ctx.db, "Study", text, uri)


def current(ctx, document_id):
    doc = ctx.db.get(Document, document_id)
    return version_chunks(ctx.db, doc.id, doc.current_version)


def test_pdf_upload_preserves_exact_raw_bytes_and_checksum(evidence_context, monkeypatch):
    ctx = evidence_context
    app = FastAPI()
    app.include_router(document_api.router)
    app.dependency_overrides[get_db] = lambda: ctx.db
    monkeypatch.setattr(document_api, "DocumentIngestion", lambda: ctx.ingestion)
    data = pdf_bytes("Method Alpha TARGETS Task Omega.")
    with TestClient(app) as client:
        response = client.post("/api/documents/import/file", files={"file": ("study.pdf", data, "application/pdf")})
    assert response.json()["code"] == 0, response.json()
    version = versions(ctx, response.json()["data"]["document_id"])[0]
    assert ctx.storage.load(version.source_storage_key) == data
    assert version.source_checksum == hashlib.sha256(data).hexdigest()


def test_identical_pdf_upload_deduplicates_blob_and_version(evidence_context):
    ctx = evidence_context
    data = pdf_bytes("Method Alpha TARGETS Task Omega.")
    first = ctx.ingestion.import_bytes(ctx.db, data, "pdf", "study.pdf")
    second = ctx.ingestion.import_bytes(ctx.db, data, "pdf", "study.pdf")
    assert second["index_action"] == "unchanged" and second["version"] == 1
    assert len(versions(ctx, first["document_id"])) == 1
    assert len(list(ctx.storage.iter_keys())) == 1


def test_same_filename_different_bytes_preserves_both_snapshots(evidence_context):
    ctx = evidence_context
    first = ctx.ingestion.import_bytes(ctx.db, pdf_bytes("First source."), "pdf", "study.pdf")
    ctx.ingestion.import_bytes(ctx.db, pdf_bytes("Changed source."), "pdf", "study.pdf")
    rows = versions(ctx, first["document_id"])
    assert len(rows) == 2 and rows[0].source_checksum != rows[1].source_checksum
    assert all(ctx.storage.exists(v.source_storage_key) for v in rows)


def test_duplicate_pdf_chunks_have_distinct_occurrences_and_pages(evidence_context):
    ctx = evidence_context
    result = ctx.ingestion.import_bytes(ctx.db, pdf_bytes("Identical paragraph.", "Identical paragraph."), "pdf", "repeat.pdf")
    rows = current(ctx, result["document_id"])
    assert len(rows) == 2 and rows[0].chunk_hash == rows[1].chunk_hash
    assert rows[0].stable_chunk_id != rows[1].stable_chunk_id
    assert [c.occurrence_index for c in rows] == [0, 1]
    assert [c.page_number for c in rows] == [1, 2]
    assert all(resolve_version_chunk(ctx.db, c.document_version_id, c.stable_chunk_id) for c in rows)


def test_update_reuses_ab_and_embeds_only_d(evidence_context, monkeypatch):
    ctx = evidence_context
    result = text_import(ctx)
    before = current(ctx, result["document_id"])
    spy = Mock(wraps=ctx.store.embedding.embed_documents)
    monkeypatch.setattr(ctx.store.embedding, "embed_documents", spy)
    updated = text_import(ctx, "Alpha evidence.\n\nBeta evidence.\n\nNew evidence.")
    after = current(ctx, result["document_id"])
    assert [c.id for c in after[:2]] == [c.id for c in before[:2]]
    assert after[2].id != before[2].id
    assert updated["stats"]["reembedded_chunk_count"] == 1
    assert [text for call in spy.call_args_list for text in call.args[0]] == ["New evidence."]
    assert ctx.db.get(DocumentChunk, before[2].id) is not None


def test_old_citation_resolves_after_update(evidence_context):
    ctx = evidence_context
    result = text_import(ctx)
    old = current(ctx, result["document_id"])[2]
    expected = resolve_version_chunk(ctx.db, old.document_version_id, old.stable_chunk_id)
    text_import(ctx, "Alpha evidence.\n\nBeta evidence.\n\nNew evidence.")
    assert resolve_version_chunk(ctx.db, old.document_version_id, old.stable_chunk_id) == expected


def test_retained_occurrence_location_is_a_version_snapshot(evidence_context):
    ctx = evidence_context
    first = ParsedDocument("Locations", "text", "phase2:locations", "Alpha\nBeta",
                           [ParsedSection("Alpha", 0, 1, "Old section"), ParsedSection("Beta", 1, 2, "Other")])
    result = ctx.indexer.import_parsed(ctx.db, first)
    old = current(ctx, result["document_id"])[0]
    old_reference = resolve_version_chunk(ctx.db, old.document_version_id, old.id)
    # Different parsed evidence/config creates a new version while Alpha is retained.
    get_settings().parser_version = "parser-v2"
    newer = replace(first, sections=[ParsedSection("Alpha", 0, 9, "New section"), ParsedSection("Beta", 1, 10, "Other")])
    ctx.indexer.import_parsed(ctx.db, newer)
    new = current(ctx, result["document_id"])[0]
    assert new.id == old.id and new.page_number == 9 and new.section_title == "New section"
    assert resolve_version_chunk(ctx.db, old.document_version_id, old.id) == old_reference


@pytest.mark.parametrize("mode", ["vector", "hybrid", "graph_enhanced"])
def test_search_excludes_historical_only_sql_and_vectors(evidence_context, mode):
    ctx = evidence_context
    result = text_import(ctx)
    old = current(ctx, result["document_id"])[2]
    # Explicit Graph now requires real current-ready Graph enrichment.
    text_import(ctx, "Method Alpha USES Dataset Beta. Alpha evidence.\n\nBeta evidence.\n\nNew evidence.")
    assert ctx.db.get(DocumentChunk, old.id) is not None
    assert old.stable_chunk_id not in ctx.store.metadata_for_ids([old.stable_chunk_id])
    # Even manually reintroduced historical vectors are not admissible.
    ctx.store.upsert_document_chunks([old])
    hits = ResearchRetrievalService(ctx.db, ctx.store).search("Old evidence", mode=mode, top_k=10)["results"]
    assert old.stable_chunk_id not in [hit["chunk_id"] for hit in hits]
    assert all(hit["document"]["version"] == 2 for hit in hits)


@pytest.mark.parametrize("mode", ["embedding_only", "graph_only"])
def test_rebuild_does_not_change_versions_or_evidence(evidence_context, mode):
    ctx = evidence_context
    result = text_import(ctx)
    before = current(ctx, result["document_id"])
    frozen = [resolve_version_chunk(ctx.db, c.document_version_id, c.id) for c in before]
    ctx.indexer.reindex(ctx.db, ctx.db.get(Document, result["document_id"]), mode)
    assert len(versions(ctx, result["document_id"])) == 1
    assert [resolve_version_chunk(ctx.db, c.document_version_id, c.id) for c in before] == frozen


def test_parser_change_reprocess_creates_version_from_same_raw(evidence_context):
    ctx = evidence_context
    result = text_import(ctx)
    old = versions(ctx, result["document_id"])[0]
    get_settings().parser_version = "parser-v2"
    result = ctx.ingestion.reprocess(ctx.db, ctx.db.get(Document, result["document_id"]))
    rows = versions(ctx, result["document_id"])
    assert result["version"] == 2 and rows[1].source_checksum == old.source_checksum
    assert rows[0].parser_version != rows[1].parser_version


def test_chunk_strategy_change_reprocess_creates_new_membership(evidence_context):
    ctx = evidence_context
    result = text_import(ctx, "A long research paragraph. " * 20)
    old = versions(ctx, result["document_id"])[0]
    get_settings().chunk_size = 100
    get_settings().chunk_overlap = 20
    result = ctx.ingestion.reprocess(ctx.db, ctx.db.get(Document, result["document_id"]))
    new = versions(ctx, result["document_id"])[1]
    assert result["version"] == 2 and new.source_checksum == old.source_checksum
    assert new.chunking_config_hash != old.chunking_config_hash
    assert len(current(ctx, result["document_id"])) > 1


def test_legacy_search_works_but_reprocess_is_explicitly_unavailable(evidence_context):
    ctx = evidence_context
    parsed = ParsedDocument("Legacy", "text", "legacy", "Legacy evidence.", [ParsedSection("Legacy evidence.", 0)])
    result = ctx.indexer.import_parsed(ctx.db, parsed)
    doc = ctx.db.get(Document, result["document_id"])
    assert ResearchRetrievalService(ctx.db, ctx.store).search("Legacy", mode="vector")["results"]
    with pytest.raises(RawSourceUnavailable, match="raw_source_unavailable"):
        ctx.ingestion.reprocess(ctx.db, doc)


def test_url_reprocess_uses_snapshot_without_refetch(evidence_context, monkeypatch):
    ctx = evidence_context
    data = b"<html><title>Research</title><p>Original evidence.</p></html>"
    fetch = Mock(return_value={"data": data, "content_type": "text/html", "encoding": "utf-8"})
    monkeypatch.setattr(UrlParser, "fetch_snapshot", fetch)
    result = ctx.ingestion.import_url(ctx.db, "https://example.test/study")
    fetch.side_effect = AssertionError("Must not fetch changed server content on reprocess")
    get_settings().parser_version = "parser-v2"
    ctx.ingestion.reprocess(ctx.db, ctx.db.get(Document, result["document_id"]))
    assert fetch.call_count == 1
    assert "Original evidence." in " ".join(c.text for c in current(ctx, result["document_id"]))
    assert all(ctx.storage.load(v.source_storage_key) == data for v in versions(ctx, result["document_id"]))


def test_traversal_filename_never_controls_blob_path(evidence_context):
    ctx = evidence_context
    result = ctx.ingestion.import_bytes(ctx.db, pdf_bytes("Safe evidence."), "pdf", "../../secret.pdf")
    version = versions(ctx, result["document_id"])[0]
    assert version.source_storage_key.startswith("sources/")
    assert ".." not in version.source_storage_key and "secret" not in version.source_storage_key
    assert ctx.storage.load(version.source_storage_key).startswith(b"%PDF")


def test_shared_source_is_deleted_only_after_last_document(evidence_context):
    ctx = evidence_context
    first = text_import(ctx, "Shared bytes.", "source:one")
    second = text_import(ctx, "Shared bytes.", "source:two")
    assert first["document_id"] != second["document_id"]
    key = versions(ctx, first["document_id"])[0].source_storage_key
    assert len(list(ctx.storage.iter_keys())) == 1
    ctx.indexer.delete_document(ctx.db, ctx.db.get(Document, first["document_id"]))
    assert ctx.storage.exists(key)
    ctx.indexer.delete_document(ctx.db, ctx.db.get(Document, second["document_id"]))
    assert not ctx.storage.exists(key)


def test_failed_sql_prepare_leaves_only_collectable_source(evidence_context, monkeypatch):
    ctx = evidence_context
    monkeypatch.setattr(ctx.indexer, "_create_version", Mock(side_effect=RuntimeError("SQL failure")))
    with pytest.raises(IncrementalIndexError):
        text_import(ctx)
    assert ctx.db.scalar(select(func.count(DocumentVersion.id))) == 0
    gc = SourceGarbageCollector(ctx.db, ctx.storage)
    assert gc.run()["orphan_count"] == 1
    assert len(list(ctx.storage.iter_keys())) == 1
    assert gc.run(repair=True)["deleted_count"] == 1
    assert gc.run(repair=True)["deleted_count"] == 0


def test_reconciliation_uses_current_membership_not_all_historical_sql(evidence_context):
    ctx = evidence_context
    result = text_import(ctx)
    old = current(ctx, result["document_id"])[2]
    text_import(ctx, "Alpha evidence.\n\nBeta evidence.\n\nNew evidence.")
    active = current(ctx, result["document_id"])
    ctx.store.upsert_document_chunks([old])
    ctx.store.delete_chunk_ids([active[0].stable_chunk_id])
    ctx.store.collection.update(ids=[active[1].stable_chunk_id], metadatas=[{"document_version_id": 99999}])
    service = VectorReconciliationService(ctx.db, ctx.store, batch_size=2)
    report = service.run()
    assert report["orphan_vector_ids"] == [old.stable_chunk_id]
    assert report["missing_vector_ids"] == [active[0].stable_chunk_id]
    assert report["stale_vector_ids"] == [active[1].stable_chunk_id]
    service.run(repair=True)
    assert service.run()["orphan_count"] == 0 and service.run()["missing_count"] == 0
    assert ctx.db.get(DocumentChunk, old.id) is not None


def test_first_citation_survives_multiple_updates_and_reprocesses(evidence_context):
    ctx = evidence_context
    result = text_import(ctx)
    row = current(ctx, result["document_id"])[2]
    expected = resolve_version_chunk(ctx.db, row.document_version_id, row.id)
    text_import(ctx, "Alpha evidence.\n\nNew evidence.")
    get_settings().parser_version = "parser-v2"
    ctx.ingestion.reprocess(ctx.db, ctx.db.get(Document, result["document_id"]))
    get_settings().chunk_size = 300
    get_settings().chunk_overlap = 40
    ctx.ingestion.reprocess(ctx.db, ctx.db.get(Document, result["document_id"]))
    assert len(versions(ctx, result["document_id"])) == 4
    assert resolve_version_chunk(ctx.db, row.document_version_id, row.id) == expected


def test_failed_new_version_does_not_switch_current(evidence_context, monkeypatch):
    ctx = evidence_context
    result = text_import(ctx)
    publish = ctx.store.upsert_document_chunks
    monkeypatch.setattr(ctx.store, "upsert_document_chunks", Mock(side_effect=RuntimeError("publish failed")))
    with pytest.raises(IncrementalIndexError):
        text_import(ctx, "Changed evidence.")
    doc = ctx.db.get(Document, result["document_id"])
    assert doc.current_version == 1 and doc.pending_version == 2 and doc.status == "publish_failed"
    monkeypatch.setattr(ctx.store, "upsert_document_chunks", publish)
    ctx.indexer.reindex(ctx.db, doc, "embedding_only")
    assert doc.current_version == 2 and doc.pending_version is None
    assert len(versions(ctx, doc.id)) == 2


def test_version_history_and_reprocess_api_contract(evidence_context, monkeypatch):
    ctx = evidence_context
    result = text_import(ctx)
    app = FastAPI()
    app.include_router(document_api.router)
    app.dependency_overrides[get_db] = lambda: ctx.db
    monkeypatch.setattr(document_api, "DocumentIngestion", lambda: ctx.ingestion)
    with TestClient(app) as client:
        history = client.get(f"/api/documents/{result['document_id']}/versions").json()
        assert history["code"] == 0 and history["data"][0]["raw_source_available"]
        get_settings().parser_version = "parser-v2"
        response = client.post(f"/api/documents/{result['document_id']}/reprocess", json={}).json()
        assert response["code"] == 0 and response["data"]["version"] == 2
        assert client.get(f"/api/documents/{result['document_id']}").json()["data"]["current_version"] == 2

def test_docx_bytes_are_preserved_and_reprocessable(evidence_context):
    from docx import Document as DocxDocument
    ctx = evidence_context
    word = DocxDocument()
    word.add_paragraph("Original DOCX evidence.")
    stream = io.BytesIO()
    word.save(stream)
    data = stream.getvalue()
    result = ctx.ingestion.import_bytes(ctx.db, data, "docx", "study.docx")
    version = versions(ctx, result["document_id"])[0]
    assert ctx.storage.load(version.source_storage_key) == data
    get_settings().parser_version = "parser-v2"
    assert ctx.ingestion.reprocess(ctx.db, ctx.db.get(Document, result["document_id"]))["version"] == 2


def test_existing_version_and_membership_reject_in_place_changes(evidence_context):
    ctx = evidence_context
    result = text_import(ctx)
    version = versions(ctx, result["document_id"])[0]
    version.parser_version = "illegal-mutation"
    with pytest.raises(ValueError, match="immutable"):
        ctx.db.commit()
    ctx.db.rollback()
    membership = ctx.db.scalar(select(DocumentVersionChunk).where(DocumentVersionChunk.document_version_id == version.id))
    membership.page_number = 999
    with pytest.raises(ValueError, match="immutable"):
        ctx.db.commit()
    ctx.db.rollback()


def test_legacy_reupload_creates_a_real_source_version(evidence_context):
    ctx = evidence_context
    parsed = ParsedDocument("Study", "text", "phase2:study", "Legacy evidence.", [ParsedSection("Legacy evidence.", 0)])
    first = ctx.indexer.import_parsed(ctx.db, parsed)
    text_import(ctx, "Legacy evidence.")
    rows = versions(ctx, first["document_id"])
    assert len(rows) == 2
    assert rows[0].source_storage_key is None and rows[1].source_storage_key is not None

def test_source_lock_failure_does_not_block_sql_delete(evidence_context, monkeypatch):
    from contextlib import contextmanager
    ctx = evidence_context
    result = text_import(ctx)
    key = versions(ctx, result["document_id"])[0].source_storage_key
    original = LocalSourceStorage.mutation_lock
    @contextmanager
    def unavailable(self):
        raise OSError("storage unavailable")
        yield
    monkeypatch.setattr(LocalSourceStorage, "mutation_lock", unavailable)
    deleted = ctx.indexer.delete_document(ctx.db, ctx.db.get(Document, result["document_id"]))
    assert deleted["cleanup_pending"]
    assert ctx.db.get(Document, result["document_id"]) is None
    assert ctx.storage.exists(key)
    monkeypatch.setattr(LocalSourceStorage, "mutation_lock", original)
    assert SourceGarbageCollector(ctx.db, ctx.storage).run(repair=True)["deleted_count"] == 1
