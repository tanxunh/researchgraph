from __future__ import annotations

import io
from pathlib import Path

from docx import Document as DocxDocument

from app.services.parsing.base import ParsedDocument, ParsedSection, ParserError


class DocxParser:
    source_type = "docx"

    def parse(self, filename: str, data: bytes) -> ParsedDocument:
        if not data:
            raise ParserError("Uploaded DOCX is empty.")
        try:
            docx = DocxDocument(io.BytesIO(data))
        except Exception as exc:
            raise ParserError(f"DOCX could not be opened: {exc}") from exc

        sections: list[ParsedSection] = []
        current_title: str | None = None
        order = 0
        texts: list[str] = []
        for paragraph in docx.paragraphs:
            text = paragraph.text.strip()
            if not text:
                continue
            style_name = (paragraph.style.name or "").lower() if paragraph.style else ""
            if style_name.startswith("heading"):
                current_title = text[:255]
            texts.append(text)
            sections.append(ParsedSection(text=text, order=order, section_title=current_title, source_uri=filename))
            order += 1

        full_text = "\n\n".join(texts).strip()
        if not full_text:
            raise ParserError("DOCX text extraction returned empty content.")
        document = ParsedDocument(
            title=(Path(filename).stem or "DOCX Document")[:255],
            source_type="docx",
            source_uri=filename,
            text=full_text,
            sections=sections,
            metadata={"filename": filename},
        )
        document.validate()
        return document
